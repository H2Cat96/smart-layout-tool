#!/usr/bin/env python3
"""Extract PDF text into a DOCX source for the gonggu-neiye template.

This is a pragmatic importer for text-layer PDFs. It preserves answer blanks as
template answer-line sentinels and crops detected complex table regions as images
when PyMuPDF is available.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader

ANSWER_LINE_SENTINEL = "\uE000_ANSWER_LINE"

PINYIN_RE = re.compile(r"^[a-züāáǎàēéěèīíǐìōóǒòūúǔùǖǘǚǜńňḿǹ\s]+$", re.I)
PAGE_RE = re.compile(r"^·\d+·$")
LESSON_RE = re.compile(r"^第[一二三四五六七八九十百]+讲$")
SECTION_RE = re.compile(r"^(?:【)?练习[一二三四五六七八九十百]+(?:】)?$")
SOURCE_RE = re.compile(r"^[（(].*(?:改编|原创|\d{4}|校考|期末|期中|测试|考试|小升初|中考).*[）)]$")
QUESTION_RE = re.compile(r"^[1-9]\d*[.．、]\s*")
OPTION_RE = re.compile(r"^[A-E][.．]\s*")
ANSWER_LABEL_RE = re.compile(r"^【?(?:答案|解析|参考答案)】?$")
CIRCLED_RE = re.compile(r"^[①②③④⑤⑥⑦⑧⑨⑩]")
PAREN_ITEM_RE = re.compile(r"^[（(]\d+[）)]")
ROMAN_HEADING_RE = re.compile(r"^[一二三四五六七八九十]、")
TABLE_HINT_RE = re.compile(r"(图表|表格|补充完整|将图表|表中)")
CJK_RE = r"\u3400-\u4dbf\u4e00-\u9fff"
CJK_PUNCT_RE = r"，。！？；：、"


def underline_for_spaces(match: re.Match[str]) -> str:
    count = len(match.group(0))
    return "_" * max(6, min(24, count // 2))


def normalize_spaced_underscores(text: str) -> str:
    def replace(match: re.Match[str]) -> str:
        underscores = match.group(0).count("_")
        return "_" * max(2, underscores)

    return re.sub(r"_(?:\s+_)+", replace, text)


def normalize_line(raw: str) -> str:
    if raw and not raw.strip() and len(raw) >= 12:
        return ANSWER_LINE_SENTINEL
    line = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", raw)
    line = line.replace("\u3000", " ").replace("\xa0", " ")
    line = re.sub(r" {4,}", underline_for_spaces, line)
    line = re.sub(r"\s+", " ", line).strip()
    line = normalize_spaced_underscores(line)
    line = re.sub(fr"(?<=[{CJK_RE}]) (?=[{CJK_RE}])", "", line)
    line = re.sub(fr"(?<=[{CJK_RE}”》）]) (?=[{CJK_PUNCT_RE}])", "", line)
    line = re.sub(fr"(?<=[{CJK_PUNCT_RE}]) (?=[{CJK_RE}“《（])", "", line)
    line = re.sub(fr"(?<=[{CJK_RE}]) (?=[“《（])", "", line)
    line = re.sub(fr"(?<=[{CJK_RE}{CJK_PUNCT_RE}]) (?=[”》）])", "", line)
    line = re.sub(fr"(?<=[”》）]) (?=[{CJK_RE}])", "", line)
    return line.replace("【参考答案】", "【答案】")


def keep_line(line: str) -> bool:
    if line == ANSWER_LINE_SENTINEL:
        return True
    if not line or PAGE_RE.match(line):
        return False
    if line == "题目与出处":
        return False
    if PINYIN_RE.match(line) and len(line.replace(" ", "")) <= 8:
        return False
    if re.fullmatch(r"\d", line):
        return False
    return True


def is_boundary(line: str) -> bool:
    if line == ANSWER_LINE_SENTINEL:
        return True
    if LESSON_RE.match(line) or SECTION_RE.match(line) or ANSWER_LABEL_RE.match(line):
        return True
    if SOURCE_RE.match(line) or QUESTION_RE.match(line) or OPTION_RE.match(line):
        return True
    if CIRCLED_RE.match(line) or PAREN_ITEM_RE.match(line) or ROMAN_HEADING_RE.match(line):
        return True
    return line in {"参考答案", "参考答案与解析"}


def should_join(prev: str, line: str) -> bool:
    if not prev or prev == ANSWER_LINE_SENTINEL or line == ANSWER_LINE_SENTINEL:
        return False
    if is_boundary(line):
        return False
    if LESSON_RE.match(prev) or SECTION_RE.match(prev) or ANSWER_LABEL_RE.match(prev):
        return False
    if CIRCLED_RE.match(prev):
        return True
    if prev.endswith(("。", "！", "？", "”", "）", ")", "：", ":")) and not line.startswith(("，", "。", "；", "、")):
        return False
    return True


def compact_answer_lines(blocks: list[Any]) -> list[Any]:
    compact: list[Any] = []
    answer_line_run = 0
    for block in blocks:
        if block == ANSWER_LINE_SENTINEL:
            answer_line_run += 1
            if answer_line_run <= 4:
                compact.append(block)
        else:
            answer_line_run = 0
            compact.append(block)
    return compact


def expand_table_union_with_connectors(union: Any, drawings: list[dict[str, Any]]) -> Any:
    expanded = union
    for drawing in drawings:
        rect = drawing.get("rect")
        if not rect:
            continue
        # Include narrow vertical connector lines that hang from flowchart/table
        # boxes. They are often separate path rectangles below the main cell row.
        narrow = rect.width <= 8
        horizontally_inside = union.x0 - 8 <= rect.x0 <= union.x1 + 8 or union.x0 - 8 <= rect.x1 <= union.x1 + 8
        starts_near_bottom = union.y0 <= rect.y0 <= union.y1 + 12
        extends_down = rect.y1 > union.y1
        if narrow and horizontally_inside and starts_near_bottom and extends_down:
            expanded = expanded | rect
    return expanded


def should_crop_drawing_cluster(cluster: list[Any], union: Any) -> bool:
    if union.width < 180:
        return False
    if union.height >= 70:
        return True
    if len(cluster) >= 4 and union.height >= 35:
        return True
    return False


def rects_intersect(a: Any, b: Any) -> bool:
    return not (a.x1 <= b.x0 or a.x0 >= b.x1 or a.y1 <= b.y0 or a.y0 >= b.y1)


def detect_table_crops(pdf_path: Path, image_dir: Path) -> dict[int, list[dict[str, Any]]]:
    try:
        import fitz  # type: ignore[import-not-found]
    except ModuleNotFoundError:
        return {}

    image_dir.mkdir(parents=True, exist_ok=True)
    crops_by_page: dict[int, list[dict[str, Any]]] = {}
    doc = fitz.open(pdf_path)
    for page_index, page in enumerate(doc):
        drawings = page.get_drawings()
        rects = []
        for drawing in drawings:
            rect = drawing.get("rect")
            if not rect:
                continue
            if rect.x0 <= 40:
                continue
            if rect.width < 20 or rect.height < 8:
                continue
            # Red question badges are small rounded rectangles near the left edge;
            # keep them as text/badge styling, not screenshot figures.
            if rect.width < 40 and rect.height < 24 and rect.x0 < 90:
                continue
            if rect.width >= 20 and rect.height >= 8:
                rects.append(rect)
        clusters: list[list[Any]] = []
        for rect in rects:
            placed = False
            for cluster in clusters:
                union = cluster[0]
                if not (rect.y0 > union.y1 + 18 or rect.y1 < union.y0 - 18):
                    cluster.append(rect)
                    cluster[0] = union | rect
                    placed = True
                    break
            if not placed:
                clusters.append([rect])
        for crop_index, cluster in enumerate(clusters, 1):
            union = cluster[0]
            # Real table/chart/dictionary regions may have either many cells or a
            # single large containing frame. Red section decorations are shorter.
            if not should_crop_drawing_cluster(cluster, union):
                continue
            union = expand_table_union_with_connectors(union, drawings)
            clip = fitz.Rect(
                max(0, union.x0 - 8),
                max(0, union.y0 - 8),
                min(page.rect.x1, union.x1 + 8),
                min(page.rect.y1, union.y1 + 32),
            )
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), clip=clip, alpha=False)
            out = image_dir / f"{pdf_path.stem}_p{page_index + 1:03d}_table_{crop_index}.png"
            pix.save(out)
            crops_by_page.setdefault(page_index, []).append({"type": "image_path", "path": out, "rect": clip})
    return crops_by_page


def extract_pdf_blocks_with_positions(pdf_path: Path, table_crops: dict[int, list[dict[str, Any]]]) -> list[Any] | None:
    try:
        import fitz  # type: ignore[import-not-found]
    except ModuleNotFoundError:
        return None

    doc = fitz.open(pdf_path)
    blocks: list[Any] = []
    current = ""
    for page_index, page in enumerate(doc):
        crops = sorted(table_crops.get(page_index, []), key=lambda crop: crop["rect"].y0)
        crop_idx = 0
        text_blocks = sorted(page.get_text("blocks"), key=lambda block: (block[1], block[0]))
        page_blocks: list[Any] = []
        for text_block in text_blocks:
            x0, y0, x1, y1, text, *_ = text_block
            block_rect = fitz.Rect(x0, y0, x1, y1)
            while crop_idx < len(crops) and crops[crop_idx]["rect"].y0 <= y0:
                page_blocks.append({"type": "image_path", "path": crops[crop_idx]["path"]})
                crop_idx += 1
            if any(rects_intersect(block_rect, crop["rect"]) for crop in crops):
                continue
            for raw in text.splitlines():
                line = normalize_line(raw)
                if not keep_line(line):
                    continue
                if should_join(current, line):
                    current += line
                else:
                    if current:
                        page_blocks.append(current)
                    current = line
        while crop_idx < len(crops):
            page_blocks.append({"type": "image_path", "path": crops[crop_idx]["path"]})
            crop_idx += 1
        blocks.extend(page_blocks)
    if current:
        blocks.append(current)
    return compact_answer_lines(blocks)


def extract_pdf_blocks(pdf_path: Path, table_crops: dict[int, list[dict[str, Any]]] | None = None) -> list[Any]:
    table_crops = table_crops or {}
    positioned = extract_pdf_blocks_with_positions(pdf_path, table_crops) if table_crops else None
    if positioned is not None:
        return positioned

    reader = PdfReader(str(pdf_path))
    blocks: list[Any] = []
    current = ""
    for page_index, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if not text.strip():
            continue
        page_blocks: list[Any] = []
        for raw in text.splitlines():
            line = normalize_line(raw)
            if not keep_line(line):
                continue
            if should_join(current, line):
                current += line
            else:
                if current:
                    page_blocks.append(current)
                current = line
        crops = table_crops.get(page_index, [])
        if crops:
            inserted = False
            next_page_blocks: list[Any] = []
            for block in page_blocks:
                next_page_blocks.append(block)
                if isinstance(block, str) and TABLE_HINT_RE.search(block):
                    next_page_blocks.extend({"type": "image_path", "path": crop["path"]} for crop in crops)
                    inserted = True
            if not inserted:
                next_page_blocks.extend({"type": "image_path", "path": crop["path"]} for crop in crops)
            page_blocks = next_page_blocks
        blocks.extend(page_blocks)
    if current:
        blocks.append(current)
    return compact_answer_lines(blocks)


def write_docx(main_blocks: list[Any], answer_blocks: list[Any], output_docx: Path) -> None:
    doc = Document()
    for block in main_blocks:
        if isinstance(block, dict) and block.get("type") == "image_path":
            doc.add_picture(str(block["path"]))
        else:
            doc.add_paragraph(str(block))
    doc.add_paragraph("参考答案与解析")
    for block in answer_blocks:
        if isinstance(block, dict) and block.get("type") == "image_path":
            doc.add_picture(str(block["path"]))
        else:
            doc.add_paragraph(str(block))
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_docx)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--main-pdf", required=True)
    parser.add_argument("--answer-pdf", required=True)
    parser.add_argument("--out-docx", required=True)
    parser.add_argument("--image-dir", required=True)
    args = parser.parse_args()

    main_pdf = Path(args.main_pdf)
    answer_pdf = Path(args.answer_pdf)
    image_dir = Path(args.image_dir)
    main_crops = detect_table_crops(main_pdf, image_dir)
    answer_crops = detect_table_crops(answer_pdf, image_dir / "answers")
    main_blocks = extract_pdf_blocks(main_pdf, main_crops)
    answer_blocks = [block for block in extract_pdf_blocks(answer_pdf, answer_crops) if block != "参考答案"]
    write_docx(main_blocks, answer_blocks, Path(args.out_docx))
    print({
        "out_docx": args.out_docx,
        "main_blocks": len(main_blocks),
        "answer_blocks": len(answer_blocks),
        "main_table_crops": sum(len(v) for v in main_crops.values()),
        "answer_table_crops": sum(len(v) for v in answer_crops.values()),
    })


if __name__ == "__main__":
    main()
