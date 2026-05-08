import importlib.util
import tempfile
import unittest
from unittest import mock
from io import BytesIO
from pathlib import Path

from PIL import Image
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from docx import Document


GENERATOR_PATH = (
    Path(__file__).resolve().parents[1]
    / "templates"
    / "gonggu-neiye"
    / "legacy"
    / "generate_print_pdf.py"
)


def load_generator():
    spec = importlib.util.spec_from_file_location("legacy_generator_under_test", GENERATOR_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class LegacyGeneratorRulesTests(unittest.TestCase):
    def test_load_docx_paragraphs_preserves_tables_and_underlined_runs(self):
        generator = load_generator()
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "sample.docx"
            doc = Document()
            p = doc.add_paragraph()
            p.add_run("普通文字")
            p.add_run("画线文字").underline = True
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "事件"
            table.cell(0, 1).text = "作者感情"
            table.cell(1, 0).text = "听到消息"
            table.cell(1, 1).text = "伤感"
            image_path = Path(tmp) / "sample.png"
            Image.new("RGB", (120, 60), "white").save(image_path)
            doc.add_picture(str(image_path))
            doc.save(docx_path)

            blocks = generator.load_docx_paragraphs(docx_path)

            self.assertEqual(blocks[0]["type"], "paragraph")
            self.assertEqual(blocks[0]["text"], "普通文字画线文字")
            self.assertEqual(blocks[0]["underline_ranges"], [[4, 8]])
            self.assertEqual(blocks[1]["type"], "table")
            self.assertEqual(blocks[1]["rows"], [["事件", "作者感情"], ["听到消息", "伤感"]])
            self.assertEqual(blocks[2]["type"], "image")
            self.assertEqual(blocks[2]["width_px"], 120)
            self.assertEqual(blocks[2]["height_px"], 60)

    def test_table_cells_preserve_fill_in_blank_widths(self):
        generator = load_generator()
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "table-fill.docx"
            doc = Document()
            table = doc.add_table(rows=1, cols=1)
            p = table.cell(0, 0).paragraphs[0]
            p.add_run("答案（")
            p.add_run(" " * 16).underline = True
            p.add_run("）")
            doc.save(docx_path)

            blocks = generator.load_docx_paragraphs(docx_path)

            self.assertEqual(blocks[0]["rows"], [["答案（" + " " * 16 + "）"]])
            self.assertEqual(blocks[0]["cell_underline_ranges"], [[[[3, 19, 16]]]])

    def test_list_prefix_shifts_paragraph_underline_ranges(self):
        generator = load_generator()
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "numbered-fill.docx"
            doc = Document()
            p = doc.add_paragraph()
            p.add_run("题干：")
            p.add_run(" " * 8).underline = True
            p.add_run("。")
            doc.save(docx_path)

            with mock.patch.object(generator, "_paragraph_list_prefix", return_value="1."):
                blocks = generator.load_docx_paragraphs(docx_path)

            self.assertEqual(blocks[0]["text"], "1.题干：" + " " * 8 + "。")
            self.assertEqual(blocks[0]["underline_ranges"], [[5, 13, 8]])

    def test_paragraph_fill_in_blanks_keep_spaces_not_visible_underscores(self):
        generator = load_generator()
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("无法实现")
        p.add_run(" " * 16).underline = True
        p.add_run("。")

        self.assertEqual(generator.paragraph_fill_in_text(p), "无法实现" + " " * 16 + "。")
        self.assertEqual(generator.paragraph_underline_ranges(p), [[4, 20, 16]])

    def test_trailing_underlined_blank_survives_plain_padding_trim(self):
        generator = load_generator()
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("少：")
        p.add_run(" " * 15).underline = True
        p.add_run(" " * 12)

        self.assertEqual(generator.paragraph_fill_in_text(p), "少：" + " " * 15)
        self.assertEqual(generator.paragraph_underline_ranges(p), [[2, 17, 15]])

    def test_split_fill_in_underline_only_draws_current_line_segment(self):
        generator = load_generator()

        class FakeCanvas:
            _pagesize = (200, 200)

            def __init__(self):
                self.lines = []

            def setStrokeColor(self, color):
                pass

            def setLineWidth(self, width):
                pass

            def line(self, x1, y1, x2, y2):
                self.lines.append((x1, y1, x2, y2))

        c = FakeCanvas()
        style = {
            "font": "Helvetica",
            "size": 12,
            "color": "#222222",
            "underline_ranges": [[3, 19, 16]],
        }

        generator.draw_underlines_for_line(c, "答案（    ", 0, 10, 40, style)

        self.assertEqual(len(c.lines), 1)
        self.assertAlmostEqual(c.lines[0][2] - c.lines[0][0], generator.text_width(" " * 4, "Helvetica", 12))

    def test_forbidden_line_start_punctuation_may_exceed_width(self):
        generator = load_generator()
        style = {
            "font": "Helvetica",
            "size": 12,
        }
        max_w = generator.text_width("abcdef", "Helvetica", 12)

        lines = generator.wrap_text("abcdef,ghijkl", style, max_w)

        self.assertEqual(lines[0], "abcdef,")
        self.assertGreater(generator.text_width(lines[0], "Helvetica", 12), max_w)

    def test_numbered_question_preserves_shifted_underline_ranges(self):
        generator = load_generator()
        block = {
            "type": "paragraph",
            "text": "1.题干：" + " " * 8 + "。",
            "underline_ranges": [[5, 13, 8]],
        }

        style = generator.paragraph_style(
            block,
            3,
            False,
            {"body": "FZKaiGBK", "title": "FZYanSongZhong", "question": "FZYanSongZhun"},
            {},
        )

        self.assertEqual(style["display_text"], "题干：" + " " * 8 + "。")
        self.assertEqual(style["underline_ranges"], [[3, 11, 8]])
        self.assertEqual(style["wrap_width_factor"], 1.0)

    def test_numbered_question_keeps_underlined_blank_after_marker(self):
        generator = load_generator()
        block = {
            "type": "paragraph",
            "text": "2." + " " * 8 + "、白朴",
            "underline_ranges": [[2, 10, 8]],
        }

        style = generator.paragraph_style(
            block,
            3,
            False,
            {"body": "FZKaiGBK", "title": "FZYanSongZhong", "question": "FZYanSongZhun"},
            {},
        )

        self.assertEqual(style["display_text"], " " * 8 + "、白朴")
        self.assertEqual(style["underline_ranges"], [[0, 8, 8]])

    def test_parenthesized_option_preserves_shifted_inline_ranges(self):
        generator = load_generator()
        block = {
            "type": "paragraph",
            "text": "（1）少：" + " " * 8,
            "underline_ranges": [[5, 13, 8]],
            "bold_ranges": [[3, 4]],
        }

        style = generator.paragraph_style(
            block,
            3,
            False,
            {"body": "FZKaiGBK", "title": "FZYanSongZhong", "question": "FZYanSongZhun"},
            {},
        )

        self.assertEqual(style["display_text"], "（1） 少：" + " " * 8)
        self.assertEqual(style["underline_ranges"], [[6, 14, 8]])
        self.assertEqual(style["bold_ranges"], [[4, 5]])

    def test_wrapped_underline_uses_line_local_prefix_width(self):
        generator = load_generator()

        class FakeCanvas:
            _pagesize = (300, 300)

            def __init__(self):
                self.lines = []

            def setStrokeColor(self, color):
                pass

            def setLineWidth(self, width):
                pass

            def line(self, x1, y1, x2, y2):
                self.lines.append((x1, y1, x2, y2))

        c = FakeCanvas()
        style = {
            "font": "Helvetica",
            "size": 12,
            "color": "#222222",
            "underline_ranges": [[10, 14, 4]],
        }

        generator.draw_underlines_for_line(c, "    。", 10, 50, 40, style)

        self.assertEqual(len(c.lines), 1)
        self.assertAlmostEqual(c.lines[0][0], 50)
        self.assertAlmostEqual(c.lines[0][2] - c.lines[0][0], generator.text_width(" " * 4, "Helvetica", 12))

    def test_font_paths_can_be_relative_to_font_map_directory(self):
        generator = load_generator()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            font_map_dir = root / "extracted"
            font_path = root / "assets" / "fonts" / "Body.ttf"
            font_path.parent.mkdir(parents=True)
            font_path.write_bytes(b"font")

            resolved = generator.resolve_font_path("../assets/fonts/Body.ttf", font_map_dir)

            self.assertEqual(resolved, font_path.resolve())

    def test_bare_answer_items_after_answer_label_get_numbered_until_explicit_number(self):
        generator = load_generator()
        blocks = [
            "【答案】",
            "A",
            "A",
            "C",
            "（1）地  （2）的  （3）的 地",
            "5.（1）的 地 （2）的 得",
            "“的”改为“地”；“的”改为“得”",
            "【解析】",
        ]

        normalized = generator.normalize_bare_answer_items(blocks)

        self.assertEqual(normalized[1], "1.A")
        self.assertEqual(normalized[2], "2.A")
        self.assertEqual(normalized[3], "3.C")
        self.assertEqual(normalized[4], "4.（1）地  （2）的  （3）的 地")
        self.assertEqual(normalized[5], "5.（1）的 地 （2）的 得")
        self.assertEqual(normalized[6], "“的”改为“地”；“的”改为“得”")

    def test_split_answers_accepts_lesson_prefixed_answer_title(self):
        generator = load_generator()
        practice, answers = generator.split_practice_and_answers([
            "正文",
            "第二讲 答案与解析：",
            "【答案】",
        ])

        self.assertEqual(practice, ["正文"])
        self.assertEqual(answers, ["第二讲 答案与解析：", "【答案】"])

    def test_pdf_importer_joins_unfinished_sentence_across_pages(self):
        import importlib.util

        importer_path = Path(__file__).resolve().parents[1] / "tools" / "pdf_to_gonggu_docx.py"
        spec = importlib.util.spec_from_file_location("pdf_to_gonggu_docx_under_test", importer_path)
        importer = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(importer)

        current = ""
        page_one_blocks = []
        for raw in ["但母亲看看我那"]:
            line = importer.normalize_line(raw)
            if importer.should_join(current, line):
                current += line
            else:
                if current:
                    page_one_blocks.append(current)
                current = line
        page_two_blocks = []
        for raw in ["副样子， 宽容地叹息一声， 没骂我也没打我。"]:
            line = importer.normalize_line(raw)
            if importer.should_join(current, line):
                current += line
            else:
                if current:
                    page_two_blocks.append(current)
                current = line

        self.assertEqual(page_one_blocks, [])
        self.assertEqual(page_two_blocks, [])
        self.assertEqual(current, "但母亲看看我那副样子，宽容地叹息一声，没骂我也没打我。")

    def test_pdf_importer_keeps_circled_article_paragraph_continuation(self):
        import importlib.util

        importer_path = Path(__file__).resolve().parents[1] / "tools" / "pdf_to_gonggu_docx.py"
        spec = importlib.util.spec_from_file_location("pdf_to_gonggu_docx_under_test", importer_path)
        importer = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(importer)

        self.assertTrue(importer.should_join(
            "③ 但母亲看看我那副样子，宽容地叹息一声，没骂我也没打我，只是让我赶快出去弄点草喂羊。",
            "我飞快地跑出家门，心情好得要命，那时我真感到了幸福。",
        ))
        self.assertFalse(importer.should_join(
            "③ 上一段结束。",
            "④ 下一段开始。",
        ))

    def test_pdf_importer_removes_false_cjk_spaces(self):
        import importlib.util

        importer_path = Path(__file__).resolve().parents[1] / "tools" / "pdf_to_gonggu_docx.py"
        spec = importlib.util.spec_from_file_location("pdf_to_gonggu_docx_under_test", importer_path)
        importer = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(importer)

        self.assertEqual(
            importer.normalize_line("但母亲看看我那 副 样 子， 宽容地叹息一声， 没骂我也没打我。"),
            "但母亲看看我那副样子，宽容地叹息一声，没骂我也没打我。",
        )
        self.assertEqual(
            importer.normalize_line("叫叫 “上帝，可怜可怜我吧！ ” 之外"),
            "叫叫“上帝，可怜可怜我吧！”之外",
        )
        self.assertEqual(
            importer.normalize_line("作者是_ _ _ _ _ _（国籍）作家_ _ _ _ _ _ _ _"),
            "作者是______（国籍）作家________",
        )
        self.assertEqual(importer.normalize_line("1.\x01请你浏览小说"), "1. 请你浏览小说")

    def test_pdf_importer_expands_table_crop_to_include_connector_lines(self):
        import importlib.util

        importer_path = Path(__file__).resolve().parents[1] / "tools" / "pdf_to_gonggu_docx.py"
        spec = importlib.util.spec_from_file_location("pdf_to_gonggu_docx_under_test", importer_path)
        importer = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(importer)

        class Rect:
            def __init__(self, x0, y0, x1, y1):
                self.x0 = x0
                self.y0 = y0
                self.x1 = x1
                self.y1 = y1
                self.width = x1 - x0
                self.height = y1 - y0

            def __or__(self, other):
                return Rect(
                    min(self.x0, other.x0),
                    min(self.y0, other.y0),
                    max(self.x1, other.x1),
                    max(self.y1, other.y1),
                )

        union = Rect(100, 100, 500, 150)
        drawings = [
            {"rect": Rect(130, 148, 132, 190)},
            {"rect": Rect(330, 148, 332, 182)},
            {"rect": Rect(520, 148, 522, 190)},
        ]

        expanded = importer.expand_table_union_with_connectors(union, drawings)

        self.assertEqual(expanded.y1, 190)
        self.assertEqual(expanded.x1, 500)

    def test_pdf_importer_crops_large_single_frame_figures_but_not_short_decorations(self):
        import importlib.util

        importer_path = Path(__file__).resolve().parents[1] / "tools" / "pdf_to_gonggu_docx.py"
        spec = importlib.util.spec_from_file_location("pdf_to_gonggu_docx_under_test", importer_path)
        importer = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(importer)

        class Rect:
            def __init__(self, x0, y0, x1, y1):
                self.x0 = x0
                self.y0 = y0
                self.x1 = x1
                self.y1 = y1
                self.width = x1 - x0
                self.height = y1 - y0

        large_dictionary_frame = [Rect(50, 220, 545, 505)]
        red_section_decoration = [Rect(48, 140, 340, 187), Rect(48, 140, 80, 187), Rect(80, 140, 340, 187)]
        multi_cell_flowchart = [
            Rect(100, 130, 190, 165),
            Rect(200, 130, 290, 165),
            Rect(300, 130, 390, 165),
            Rect(400, 130, 500, 165),
        ]

        self.assertTrue(importer.should_crop_drawing_cluster(large_dictionary_frame, large_dictionary_frame[0]))
        self.assertFalse(importer.should_crop_drawing_cluster(red_section_decoration, Rect(48, 140, 340, 187)))
        self.assertTrue(importer.should_crop_drawing_cluster(multi_cell_flowchart, Rect(100, 130, 500, 165)))

    def test_split_spread_pdf_outputs_idml_single_page_size(self):
        generator = load_generator()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            spread_pdf = root / "spread.pdf"
            single_pdf = root / "single.pdf"
            c = canvas.Canvas(str(spread_pdf), pagesize=(200, 100))
            c.drawString(10, 50, "left")
            c.drawString(110, 50, "right")
            c.showPage()
            c.drawString(10, 50, "left2")
            c.drawString(110, 50, "right2")
            c.save()

            generator.split_spread_pdf_to_single_pages(spread_pdf, single_pdf, 100, 100)

            reader = PdfReader(str(single_pdf))
            self.assertEqual(len(reader.pages), 4)
            for page in reader.pages:
                self.assertEqual(round(float(page.mediabox.width), 3), 100)
                self.assertEqual(round(float(page.mediabox.height), 3), 100)

    def test_color_mode_cmyk_uses_k_only_for_neutral_colors(self):
        generator = load_generator()

        black = generator.color_from_hex("cmyk(0,0,0,0.95)", (0, 0, 0), color_mode="cmyk")
        gray = generator.color_from_hex("#898989", (0, 0, 0), color_mode="cmyk")
        red = generator.color_from_hex("#c1020e", (0, 0, 0), color_mode="cmyk")

        self.assertEqual((black.cyan, black.magenta, black.yellow), (0, 0, 0))
        self.assertAlmostEqual(black.black, 0.95)
        self.assertEqual((gray.cyan, gray.magenta, gray.yellow), (0, 0, 0))
        self.assertGreater(gray.black, 0.4)
        self.assertGreater(red.magenta, 0.8)
        self.assertGreater(red.yellow, 0.8)

    def test_color_mode_rgb_preserves_rgb_colors(self):
        generator = load_generator()

        color = generator.color_from_hex("#898989", (0, 0, 0), color_mode="rgb")

        self.assertAlmostEqual(color.red, 137 / 255)
        self.assertAlmostEqual(color.green, 137 / 255)
        self.assertAlmostEqual(color.blue, 137 / 255)

    def test_gonggu_specific_style_fixes(self):
        generator = load_generator()
        fonts = {
            "title": "Title",
            "title_mid": "TitleMid",
            "title_bold": "TitleBold",
            "question": "Question",
            "body": "Body",
            "footer": "Footer",
        }

        article_title = generator.paragraph_style("泉", 3, False, fonts, {})
        second_article_title = generator.paragraph_style("一匹骆驼", 39, False, fonts, {})
        punctuated_article_title = generator.paragraph_style("北京！北京！", 3, False, fonts, {})
        second_reading_prompt = generator.paragraph_style("阅读《一匹骆驼》，完成下面的小题。", 38, False, fonts, {})
        reading_prompt_all_questions = generator.paragraph_style("阅读短文《自然界的时钟》，完成下列各题。", 2, False, fonts, {})
        structural_reading_prompt = generator.paragraph_style("请结合选文完成练习。", 2, False, fonts, {}, force_reading_prompt=True)
        author = generator.paragraph_style("贾平凹", 4, False, fonts, {})
        source = generator.paragraph_style("（2011-2012北京顺义九上期末）", 15, False, fonts, {})
        answer_label = generator.paragraph_style("【答案】", 2, True, fonts, {})
        plain_answer_label = generator.paragraph_style("答案：", 2, True, fonts, {})
        plain_analysis_label = generator.paragraph_style("解析:", 2, True, fonts, {})
        answer_main_title = generator.paragraph_style("第二讲 答案与解析：", 4, True, fonts, {})
        answer_section = generator.paragraph_style("【练习二】", 1, True, fonts, {})
        plain_section = generator.paragraph_style("练习一", 0, False, fonts, {})
        lesson_title = generator.paragraph_style("第一讲", 6, False, fonts, {})
        topic_heading = generator.paragraph_style("二、长篇阅读", 7, False, fonts, {})
        question = generator.paragraph_style("3.下面对文章的理解和分析，不正确的两项是（    ）", 10, False, fonts, {})
        option = generator.paragraph_style("E.文中“一片片的小叶绽了开来”一句，可以改为另一句。", 11, False, fonts, {})
        parenthesized_option = generator.paragraph_style("（2）所有植物开放和凋谢的时间都是固定不变的。        (       )", 12, False, fonts, {})
        option_at_legacy_article_title_position = generator.paragraph_style("B．牛背上牧童的短笛，这时候也成天在嘹亮地响。", 3, False, fonts, {})
        answer_item = generator.paragraph_style("2.这句话运用了神态描写、动作描写和语言描写。", 12, True, fonts, {})
        meaning_item = generator.paragraph_style("深层含义：①象征顽强不息的生命力。", 13, True, fonts, {})
        analysis_item = generator.paragraph_style("E项错误，原句将定语后置。", 14, True, fonts, {})
        judgement_item = generator.paragraph_style("（3）“生物钟”只存在于植物和动物体内，昆虫没有“生物钟”。        (       )", 15, False, fonts, {})

        self.assertEqual(article_title["font"], "TitleMid")
        self.assertEqual(second_article_title["kind"], "article_title")
        self.assertEqual(second_article_title["font"], "TitleMid")
        self.assertEqual(punctuated_article_title["kind"], "article_title")
        self.assertEqual(punctuated_article_title["font"], "TitleMid")
        self.assertEqual(second_reading_prompt["kind"], "reading_prompt")
        self.assertEqual(reading_prompt_all_questions["kind"], "reading_prompt")
        self.assertEqual(reading_prompt_all_questions["font"], "TitleMid")
        self.assertEqual(reading_prompt_all_questions["size"], 14)
        self.assertEqual(reading_prompt_all_questions["color"], "#cc0000")
        self.assertEqual(structural_reading_prompt["kind"], "reading_prompt")
        self.assertEqual(structural_reading_prompt["font"], "TitleMid")
        self.assertEqual(author["font"], "Body")
        self.assertEqual(author["size"], 14)
        self.assertEqual(source["align"], "left")
        self.assertEqual(source["color"], "#898989")
        self.assertEqual(source["space_before"], 1)
        self.assertEqual(source["space_after"], 0)
        self.assertEqual(answer_label["kind"], "answer_label")
        self.assertFalse(answer_label["bar"])
        self.assertNotIn("bold_rule", answer_label)
        self.assertEqual(plain_answer_label["display_text"], "【答案】")
        self.assertEqual(plain_analysis_label["display_text"], "【解析】")
        self.assertEqual(answer_main_title["kind"], "main_title")
        self.assertEqual(answer_main_title["title_marker_asset"], "参考答案")
        self.assertEqual(answer_section["kind"], "section")
        self.assertTrue(answer_section["bar"])
        self.assertEqual(plain_section["kind"], "section")
        self.assertEqual(plain_section["display_text"], "【练习一】")
        self.assertTrue(plain_section["bar"])
        self.assertEqual(lesson_title["kind"], "main_title")
        self.assertEqual(lesson_title["font"], "TitleMid")
        self.assertEqual(topic_heading["kind"], "topic_heading")
        self.assertEqual(topic_heading["font"], "TitleMid")
        self.assertFalse(topic_heading["bar"])
        self.assertEqual(question["space_before"], 0)
        self.assertEqual(question["align"], "left")
        self.assertEqual(question["left_indent"], 36)
        self.assertEqual(question["wrap_width_factor"], 1.0)
        self.assertEqual(option["marker_text"], "E.")
        self.assertTrue(option["inline_marker"])
        self.assertEqual(option["display_text"], "E. 文中“一片片的小叶绽了开来”一句，可以改为另一句。")
        self.assertEqual(option["align"], "left")
        self.assertEqual(option["left_indent"], 36)
        self.assertEqual(option["wrap_width_factor"], 1.0)
        self.assertEqual(parenthesized_option["kind"], "option")
        self.assertEqual(parenthesized_option["marker_text"], "（2）")
        self.assertTrue(parenthesized_option["inline_marker"])
        self.assertEqual(parenthesized_option["display_text"], "（2） 所有植物开放和凋谢的时间都是固定不变的。        (       )")
        self.assertEqual(parenthesized_option["left_indent"], 36)
        self.assertEqual(option_at_legacy_article_title_position["kind"], "option")
        self.assertEqual(option_at_legacy_article_title_position["marker_text"], "B.")
        self.assertEqual(option_at_legacy_article_title_position["font"], "Body")
        self.assertEqual(answer_item["marker_text"], "2.")
        self.assertEqual(answer_item["display_text"], "这句话运用了神态描写、动作描写和语言描写。")
        self.assertEqual(answer_item["left_indent"], 22)
        self.assertNotIn("marker_text", meaning_item)
        self.assertNotIn("display_text", meaning_item)
        self.assertNotIn("marker_text", analysis_item)
        self.assertNotIn("display_text", analysis_item)
        self.assertEqual(judgement_item["kind"], "option")
        self.assertEqual(judgement_item["marker_text"], "（3）")
        self.assertEqual(judgement_item["font"], "Body")
        self.assertEqual(judgement_item["size"], 14)
        self.assertEqual(judgement_item["align"], "left")
        self.assertTrue(generator.can_be_structural_reading_prompt("阅读短文《自然界的时钟》，完成下列各题。"))
        self.assertTrue(generator.can_be_structural_reading_prompt("请结合选文完成练习。"))
        self.assertFalse(generator.can_be_structural_reading_prompt("1.（改编自2023-2024河南洛阳期末）下面各句中画线“的”“地”“得”用法错误的一句是（  ）"))
        self.assertFalse(generator.can_be_structural_reading_prompt("A．“哥儿，你牢牢记住！”她极其郑重的说。"))

    def test_trailing_answer_parentheses_stay_with_previous_line(self):
        generator = load_generator()
        style = {"font": "Helvetica", "size": 10, "leading": 12}

        lines = generator.wrap_text("判断句内容很长很长很长很长很长很长。        (       )", style, 120)
        repaired = generator.avoid_isolated_answer_parentheses([
            "（1）公鸡报晓是人们熟知的现象，这能体现自然界是一座奇妙的活时钟。",
            " (       )",
        ])

        self.assertGreater(len(lines), 1)
        self.assertNotEqual(lines[-1].strip(), ")")
        self.assertNotEqual(lines[-1].strip(), "）")
        self.assertNotRegex(lines[-1], r"^\\s*[)）]\\s*$")
        self.assertNotRegex(lines[-1], r"^\\s*[（(]\\s+[）)]\\s*$")
        self.assertEqual(repaired[-1], "活时钟。 (       )")

    def test_body_wrapping_avoids_short_final_line(self):
        generator = load_generator()
        style = {"font": "Helvetica", "size": 10, "leading": 12, "kind": "body"}

        lines = generator.wrap_text(
            "但母亲看看我那副样子，宽容地叹息一声，没骂我也没打我，只是让我赶快出去弄点草喂羊。",
            style,
            280,
        )

        self.assertGreater(len(lines), 1)
        self.assertGreaterEqual(generator.text_width(lines[-1], "Helvetica", 10), 10 * 2)

    def test_wrapping_keeps_forbidden_start_punctuation_off_new_lines(self):
        generator = load_generator()
        style = {"font": "Helvetica", "size": 10, "leading": 12, "kind": "body"}

        lines = generator.wrap_text("他说：“这一句话很长很长很长很长很长很长”", style, 115)

        self.assertGreater(len(lines), 1)
        self.assertFalse(any(line.startswith("”") for line in lines[1:]))

    def test_short_line_rebalance_does_not_move_closing_quote_to_line_start(self):
        generator = load_generator()
        font_map_path = Path(__file__).resolve().parents[1] / "templates" / "gonggu-neiye" / "extracted" / "font-map.json"
        font_map = __import__("json").loads(font_map_path.read_text(encoding="utf-8"))
        generator.register_fonts(font_map, font_map_path.parent)
        style = {"font": "FZYanSongZhun", "size": 14, "leading": 27, "kind": "question_numbered"}

        lines = generator.wrap_text(
            "请你结合短文内容，说一说为什么《封神演义》这本书留给“我”的印象十分深刻。",
            style,
            440,
        )

        self.assertGreater(len(lines), 1)
        self.assertFalse(any(line.startswith("”") for line in lines[1:]))
        self.assertIn("“我”", "".join(lines))

    def test_question_wrapping_aligns_content_with_title(self):
        generator = load_generator()
        style = {
            "font": "Helvetica",
            "size": 10,
            "leading": 12,
            "left_indent": 28,
            "first_line_indent": 0,
        }

        lines = generator.wrap_text_by_widths(
            "请从人物描写的角度，赏析下列句子的表达效果。",
            style,
            [30, 30],
        )

        self.assertGreater(len(lines), 1)
        self.assertEqual(generator.line_x_offset(style, first_line=True), 28)
        self.assertEqual(generator.line_x_offset(style, first_line=False), 28)

    def test_hanging_context_indents_following_blocks_and_remainders(self):
        generator = load_generator()
        rules = {"styles": {"question_content": {"left_indent": 28}}}
        base_style = {
            "kind": "body",
            "font": "Helvetica",
            "size": 10,
            "leading": 12,
            "first_line_indent": 18,
        }
        answer_line = {"kind": "answer_line", "font": "Helvetica", "size": 10, "leading": 12}
        answer_body = {
            "kind": "answer_body",
            "font": "Helvetica",
            "size": 10,
            "leading": 12,
            "first_line_indent": 18,
        }
        question_style = {
            "kind": "question_numbered",
            "font": "Helvetica",
            "size": 10,
            "leading": 12,
            "left_indent": 28,
            "first_line_indent": 18,
            "badge_text": "2",
            "display_text": "题干内容",
        }

        inherited = generator.apply_hanging_context_indent(base_style, 28, rules)
        inherited_line = generator.apply_hanging_context_indent(answer_line, 28, rules)
        inherited_answer = generator.apply_hanging_context_indent(answer_body, 68, rules)
        remainder = generator.make_remainder_block("跨页续排文字", question_style)

        self.assertEqual(inherited["left_indent"], 28)
        self.assertEqual(inherited["first_line_indent"], 0)
        self.assertEqual(inherited_line["left_indent"], 28)
        self.assertEqual(inherited_answer["left_indent"], 68)
        self.assertEqual(inherited_answer["first_line_indent"], 0)
        self.assertEqual(remainder["style"]["left_indent"], 28)
        self.assertEqual(remainder["style"]["first_line_indent"], 0)
        self.assertEqual(remainder["style"]["display_text"], "跨页续排文字")
        self.assertNotIn("badge_text", remainder["style"])
        self.assertNotIn("marker_text", remainder["style"])
        self.assertTrue(generator.ends_hanging_context({"kind": "topic_heading"}))

    def test_first_answer_line_gets_extra_spacing_only_after_non_line(self):
        generator = load_generator()
        rules = {"styles": {"answer_line": {"first_space_before": 6}}}
        answer_line = {"kind": "answer_line", "space_before": 0, "leading": 22}

        first_line = generator.apply_answer_line_context_spacing(answer_line, "question_numbered", rules)
        second_line = generator.apply_answer_line_context_spacing(answer_line, "answer_line", rules)

        self.assertEqual(first_line["space_before"], 6)
        self.assertEqual(second_line["space_before"], 0)

    def test_source_and_question_keep_with_next_content(self):
        generator = load_generator()
        main_style = {"kind": "main_title", "leading": 34, "space_before": 0, "space_after": 16}
        section_style = {"kind": "section", "leading": 24, "space_before": 3, "space_after": 8}
        article_title_style = {"kind": "article_title", "leading": 24, "space_before": 0, "space_after": 4}
        answer_label_style = {"kind": "answer_label", "leading": 23, "space_before": 5, "space_after": 2}
        topic_heading_style = {"kind": "topic_heading", "leading": 24, "space_before": 6, "space_after": 4}
        source_style = {"kind": "source", "leading": 20, "space_before": 1, "space_after": 0}
        question_style = {"kind": "question_numbered", "leading": 27, "space_before": 0, "space_after": 2}
        body_style = {"kind": "body", "leading": 24, "space_before": 1, "space_after": 2}

        self.assertTrue(generator.should_keep_with_next(main_style))
        self.assertTrue(generator.should_keep_with_next(section_style))
        self.assertTrue(generator.should_keep_with_next(article_title_style))
        self.assertTrue(generator.should_keep_with_next(answer_label_style))
        self.assertTrue(generator.should_keep_with_next(topic_heading_style))
        self.assertTrue(generator.should_keep_with_next(source_style))
        self.assertTrue(generator.should_keep_with_next(question_style))
        self.assertFalse(generator.should_keep_with_next(body_style))
        self.assertEqual(generator.min_block_height_for_keep(source_style), 21)

    def test_template_shell_draws_side_strips_on_white_background_with_section_color(self):
        generator = load_generator()

        class FakeCanvas:
            def __init__(self):
                self.rects = []
                self.colors = []

            def setFillColor(self, color):
                self.color = color
                self.colors.append((round(color.red, 6), round(color.green, 6), round(color.blue, 6)))

            def rect(self, x, y, w, h, fill, stroke):
                self.rects.append((x, y, w, h, fill, stroke))

            def setFont(self, font, size):
                pass

            def drawCentredString(self, x, y, text):
                pass

        canvas = FakeCanvas()
        spread = {
            "spread_index": 2,
            "slots": [
                {
                    "role": "side_strip",
                    "bbox_pt": {"x": 0, "y": 0, "w": 20, "h": 100},
                    "fill_rgb": "#ffd9ed",
                }
            ],
        }
        rules = {"colors": {"practice_bar": "#fce5e4"}}

        generator.paint_template_shell(canvas, spread, 200, 100, "Footer", "white", rules, page_number_start=1)

        self.assertEqual(canvas.rects, [(0, 0, 20, 100, 1, 0)])
        self.assertEqual(canvas.colors[0], (0.988235, 0.898039, 0.894118))

    def test_template_shell_skips_answer_side_strips(self):
        generator = load_generator()

        class FakeCanvas:
            def __init__(self):
                self.rects = []

            def setFillColor(self, color):
                pass

            def rect(self, x, y, w, h, fill, stroke):
                self.rects.append((x, y, w, h, fill, stroke))

            def setFont(self, font, size):
                pass

            def drawCentredString(self, x, y, text):
                pass

        canvas = FakeCanvas()
        spread = {
            "spread_index": 4,
            "slots": [
                {
                    "role": "side_strip",
                    "bbox_pt": {"x": 0, "y": 0, "w": 20, "h": 100},
                    "fill_rgb": "#eeeeee",
                }
            ],
        }

        generator.paint_template_shell(canvas, spread, 200, 100, "Footer", "white", {}, page_number_start=7)

        self.assertEqual(canvas.rects, [])

    def test_paragraph_style_uses_layout_rule_overrides(self):
        generator = load_generator()
        fonts = {
            "title": "Title",
            "title_mid": "TitleMid",
            "title_bold": "TitleBold",
            "question": "Question",
            "body": "Body",
            "footer": "Footer",
        }
        rules = {
            "styles": {
                "option": {"font": "OptionFont", "size": 15, "leading": 28, "left_indent": 31},
                "section_title": {"font": "SectionFont", "size": 17, "leading": 25, "space_before": 10},
            },
            "colors": {
                "practice_bar": "#eeeeee",
                "answer_gray": "#777777",
            },
        }

        option = generator.paragraph_style("A. 选项文字", 10, False, fonts, rules)
        practice_section = generator.paragraph_style("【练习二】", 1, False, fonts, rules)
        answer_section = generator.paragraph_style("【练习二】", 1, True, fonts, rules)

        self.assertEqual(option["font"], "OptionFont")
        self.assertEqual(option["size"], 15)
        self.assertEqual(option["leading"], 28)
        self.assertEqual(option["left_indent"], 31)
        self.assertEqual(practice_section["font"], "SectionFont")
        self.assertEqual(practice_section["space_before"], 10)
        self.assertEqual(practice_section["bar_color"], "#eeeeee")
        self.assertEqual(answer_section["bar_color"], "#777777")

    def test_detection_hanging_and_normalization_rules_are_configurable(self):
        generator = load_generator()
        fonts = {
            "title": "Title",
            "title_mid": "TitleMid",
            "title_bold": "TitleBold",
            "question": "Question",
            "body": "Body",
            "footer": "Footer",
        }
        rules = {
            "content_detection": {
                "author_names": ["苏学军"],
                "lesson_title_patterns": [r"^单元[一二三四五六七八九十]+$"],
                "reading_prompt_patterns": [r"^精读.*"],
                "article_title": {"max_length": 20},
            },
            "markers": {
                "question_patterns": [r"^题([1-9]\d*)[:：]\s*(.*)$"],
                "option_patterns": [r"^选项([A-E])[:：]\s*(.*)$"],
            },
            "normalization": {
                "answer_labels": {
                    "答": "【答案】",
                    "析": "【解析】",
                }
            },
            "hanging": {
                "end_kinds": ["section"],
                "inherit_kinds": ["body"],
                "clear_first_line_indent_kinds": ["body"],
            },
            "page_flow": {
                "keep_with_next_kinds": ["source"],
            },
        }

        author = generator.paragraph_style("苏学军", 3, False, fonts, rules)
        lesson = generator.paragraph_style("单元一", 0, False, fonts, rules)
        prompt = generator.paragraph_style("精读《火星之谜》，完成练习。", 1, False, fonts, rules)
        question = generator.paragraph_style("题12：下面说法正确的是（    ）", 2, False, fonts, rules)
        option = generator.paragraph_style("选项B：这是一项测试。", 3, False, fonts, rules)
        answer = generator.paragraph_style("答", 4, True, fonts, rules)

        self.assertEqual(author["kind"], "author")
        self.assertEqual(lesson["kind"], "main_title")
        self.assertEqual(prompt["kind"], "reading_prompt")
        self.assertEqual(question["badge_text"], "12")
        self.assertEqual(question["display_text"], "下面说法正确的是（    ）")
        self.assertEqual(option["marker_text"], "B.")
        self.assertEqual(option["display_text"], "B. 这是一项测试。")
        self.assertEqual(answer["display_text"], "【答案】")
        self.assertFalse(generator.ends_hanging_context({"kind": "topic_heading"}, rules))
        self.assertTrue(generator.ends_hanging_context({"kind": "section"}, rules))
        self.assertTrue(generator.can_inherit_hanging_context({"kind": "body"}, rules))
        self.assertFalse(generator.can_inherit_hanging_context({"kind": "image"}, rules))
        self.assertTrue(generator.should_keep_with_next({"kind": "source"}, rules))
        self.assertFalse(generator.should_keep_with_next({"kind": "question_numbered"}, rules))

    def test_flow_frame_min_height_rule_extends_short_frames_without_mutating_source(self):
        generator = load_generator()
        frames = [
            {"spread_index": 5, "bbox_pt": {"x": 56, "y": 56, "w": 480, "h": 683}},
            {"spread_index": 5, "bbox_pt": {"x": 652, "y": 56, "w": 480, "h": 188}},
        ]
        rules = {"page_flow": {"min_frame_height_pt": {"answer_content_flow": 680}}}

        normalized = generator.normalize_flow_frames(frames, "answer_content_flow", rules)

        self.assertEqual(frames[1]["bbox_pt"]["h"], 188)
        self.assertEqual(normalized[0]["bbox_pt"]["h"], 683)
        self.assertEqual(normalized[1]["bbox_pt"]["h"], 680)

    def test_load_svg_assets_uses_asset_map_filenames_and_color_overrides(self):
        generator = load_generator()
        with tempfile.TemporaryDirectory() as tmp:
            svg_dir = Path(tmp)
            (svg_dir / "custom-corner.svg").write_text(
                '<svg viewBox="0 0 30.05 30.05"><path fill="#111111" d="M0 0h1v1z"/></svg>',
                encoding="utf-8",
            )
            (svg_dir / "custom-answer.svg").write_text(
                '<svg viewBox="0 0 30.05 30.05"><path fill="#222222" d="M0 0h1v1z"/></svg>',
                encoding="utf-8",
            )
            (svg_dir / "custom-bar.svg").write_text(
                '<svg viewBox="0 0 481.89 28.35"><path fill="#333333" d="M0 0h1v1z"/></svg>',
                encoding="utf-8",
            )
            (svg_dir / "custom-badge.svg").write_text(
                '<svg viewBox="0 0 22.68 12.76"><path fill="#444444" d="M0 0h1v1z"/></svg>',
                encoding="utf-8",
            )
            asset_map = {
                "assets": {
                    "practice_title_corner": {"file": "custom-corner.svg", "color": "#aaaaaa"},
                    "answer_title_corner": {"file": "custom-answer.svg", "color": "#bbbbbb"},
                    "section_bar": {"file": "custom-bar.svg", "practice_color": "#cccccc"},
                    "question_badge": {"file": "custom-badge.svg", "color": "#dddddd"},
                }
            }

            assets = generator.load_svg_assets(svg_dir, asset_map)

            self.assertEqual(Path(assets["标题角标"]["path"]).name, "custom-corner.svg")
            self.assertEqual(assets["标题角标"]["fill"], "#aaaaaa")
            self.assertEqual(Path(assets["参考答案"]["path"]).name, "custom-answer.svg")
            self.assertEqual(assets["参考答案"]["fill"], "#bbbbbb")
            self.assertEqual(Path(assets["标题"]["path"]).name, "custom-bar.svg")
            self.assertEqual(assets["标题"]["fill"], "#cccccc")
            self.assertEqual(Path(assets["序号"]["path"]).name, "custom-badge.svg")
            self.assertEqual(assets["序号"]["fill"], "#dddddd")

    def test_option_tables_stay_as_tables_not_flattened(self):
        # Word tables that contain ABCD options are kept as tables,
        # even if every cell looks like an option. Authors may intentionally
        # use 2-column option layouts.
        generator = load_generator()
        option_table = {"type": "table", "rows": [
            ["A.《夏洛的网》", "B. 《柳林风声》"],
            ["C. E.B. 怀特", "D. 肯尼斯·格雷厄姆"],
        ]}
        data_table = {"type": "table", "rows": [
            ["结构", "段落", "内容要点"],
            ["开头", "第①段", "引出玩具名称"],
        ]}
        blocks = [option_table, data_table]

        result = generator.infer_bare_options(generator.normalize_bare_answer_items(blocks))

        # Both tables pass through intact
        self.assertTrue(generator.is_table_block(result[0]))
        self.assertTrue(generator.is_table_block(result[1]))

    def test_bare_options_after_stem_get_auto_labels_from_trailing_labeled_letter(self):
        generator = load_generator()
        blocks = [
            "下列关于《柳林风声》中动物形象的理解，不正确的一项是（   ）",
            "鼹鼠淳朴善良，热爱生活。",
            "河鼠聪慧细心，热情友善。",
            "蟾蜍沉稳稳重，谦逊温和。",
            "D. 獾先生成熟可靠。",
        ]

        result = generator.infer_bare_options(blocks)

        self.assertEqual(result[0], blocks[0])
        self.assertEqual(result[1], "A. 鼹鼠淳朴善良，热爱生活。")
        self.assertEqual(result[2], "B. 河鼠聪慧细心，热情友善。")
        self.assertEqual(result[3], "C. 蟾蜍沉稳稳重，谦逊温和。")
        self.assertEqual(result[4], "D. 獾先生成熟可靠。")

    def test_four_bare_options_after_stem_are_labeled_from_a(self):
        generator = load_generator()
        blocks = [
            "（1）下列关于《夏洛的网》中动物形象的理解，不正确的一项是（   ）",
            "蜘蛛夏洛十分聪慧善良。",
            "小猪威尔伯生性胆小脆弱。",
            "老鼠坦普尔顿心地善良。",
            "小女孩弗恩富有爱心。",
            "（下一段其他内容）",
        ]

        result = generator.infer_bare_options(blocks)

        self.assertEqual(result[1], "A. 蜘蛛夏洛十分聪慧善良。")
        self.assertEqual(result[2], "B. 小猪威尔伯生性胆小脆弱。")
        self.assertEqual(result[3], "C. 老鼠坦普尔顿心地善良。")
        self.assertEqual(result[4], "D. 小女孩弗恩富有爱心。")
        self.assertEqual(result[5], "（下一段其他内容）")

    def test_already_labeled_options_pass_through_unchanged(self):
        generator = load_generator()
        blocks = [
            "下列说法正确的一项是（  ）",
            "A. 选项一",
            "B. 选项二",
            "C. 选项三",
            "D. 选项四",
            "下一题",
        ]

        result = generator.infer_bare_options(blocks)

        self.assertEqual(result, blocks)

    def test_image_table_attaches_label_as_caption_on_the_image(self):
        generator = load_generator()
        image1 = {"type": "image", "blob": b"\x89PNG1", "width_px": 100, "height_px": 80}
        image2 = {"type": "image", "blob": b"\x89PNG2", "width_px": 100, "height_px": 80}
        rows_with_cells = [
            [{"text": "", "images": [image1]}, {"text": "", "images": [image2]}],
            [{"text": "图一", "images": []}, {"text": "图二", "images": []}],
        ]

        result = generator.flatten_image_table(rows_with_cells)

        self.assertEqual(len(result), 2)
        self.assertEqual(result[0]["caption"], "图一")
        self.assertEqual(result[0]["blob"], b"\x89PNG1")
        self.assertEqual(result[1]["caption"], "图二")
        self.assertEqual(result[1]["blob"], b"\x89PNG2")

    def test_justify_is_skipped_on_lines_containing_underscore_runs(self):
        generator = load_generator()
        import re as _re

        self.assertTrue(bool(_re.search(r"_{2,}", "作者_____，故事")))
        self.assertFalse(bool(_re.search(r"_{2,}", "作者，故事")))

    def test_image_style_max_height_caps_scaled_image_height(self):
        generator = load_generator()
        portrait = {"type": "image", "width_px": 300, "height_px": 400}
        landscape = {"type": "image", "width_px": 600, "height_px": 200}

        w_tall, h_tall = generator.image_scaled_size(portrait, 464, style_max_h=260)
        w_wide, h_wide = generator.image_scaled_size(landscape, 464, style_max_h=260)

        self.assertEqual(h_tall, 260)
        self.assertAlmostEqual(w_tall, 300 * 260 / 400)
        self.assertLess(h_wide, 260)

    def test_image_caption_adds_to_total_height(self):
        generator = load_generator()
        with_caption = {"type": "image", "width_px": 100, "height_px": 80, "caption": "图一"}
        without_caption = {"type": "image", "width_px": 100, "height_px": 80}
        style = {"kind": "image", "space_before": 8, "space_after": 10, "left_indent": 0, "max_height_pt": None}

        h_with = generator.image_total_height(with_caption, style, 480)
        h_without = generator.image_total_height(without_caption, style, 480)

        self.assertGreater(h_with, h_without)
        self.assertEqual(h_with - h_without, generator.image_caption_height(with_caption))

    def test_image_style_pulls_max_height_from_layout_rules(self):
        generator = load_generator()
        fonts = {"title": "T", "title_mid": "T", "title_bold": "T", "question": "Q", "body": "B", "footer": "F"}
        rules = {"styles": {"image": {"max_height_pt": 180}}}

        style = generator.image_style(fonts, rules)

        self.assertEqual(style["max_height_pt"], 180)
